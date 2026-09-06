from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


HERE = Path(__file__).resolve().parent
OUT = HERE / "ctp-core_Declaration_of_Competing_Interests.docx"

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)

normal = doc.styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(11)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("Declaration of Competing Interests")
run.bold = True
run.font.name = "Arial"
run.font.size = Pt(16)

doc.add_paragraph()
doc.add_paragraph(
    "Manuscript title: ctp-core: An open and reproducible Python core for "
    "gamma-variate CT-perfusion analysis and standardized visualization"
)
doc.add_paragraph("Author: Shuji Yamamoto, PhD")
doc.add_paragraph()
doc.add_paragraph(
    "The author declares that he has no known competing financial interests "
    "or personal relationships that could have appeared to influence the work "
    "reported in this paper."
)
doc.add_paragraph()
doc.add_paragraph("Date: 24 August 2026")

doc.save(OUT)
print(OUT)
