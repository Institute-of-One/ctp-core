"""Check the built manuscript against the SoftwareX Original Software Publication template.

The submission SOFTX-S-26-01589 was returned by the editorial office before review:

    "Conclusions section is missing. Your current Sections 5 and 6 are also not required
     as per the template."

That is a structural test, so it is written as one. The requirements below are read off
the official template file, softwarex-osp-template.docx, downloaded from Elsevier -- not
from memory of what SoftwareX asks for.

Run against the generated file, never against the build log::

    python check_template_conformance.py

Exit status is non-zero if any requirement fails.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document

HERE = Path(__file__).resolve().parent
DOCX = HERE / "ctp-core_SoftwareX_manuscript.docx"

#: Sections 1-5, in order, exactly as the template names them. The template states that
#: all sections are mandatory unless marked optional, and Conclusions is marked mandatory.
NUMBERED_SECTIONS = [
    "1. Motivation and significance",
    "2. Software description",
    "3. Illustrative examples",
    "4. Impact",
    "5. Conclusions",
]

#: Subsections of section 2. "Sample code snippets analysis" is optional and unused.
SUBSECTIONS = ["2.1 Software architecture", "2.2 Software functionalities"]

#: The template's code-metadata rows. The left column is to be left untouched, so the
#: labels are asserted verbatim. The submitted version carried S1-S7 rows instead, which
#: are not in this template at all.
CODE_METADATA = [
    ("C1", "Current code version"),
    ("C2", "Permanent link to code/repository used for this code version"),
    ("C3", "Legal code license"),
    ("C4", "Code versioning system used"),
    ("C5", "Software code languages, tools and services used"),
    ("C6", "Compilation requirements, operating environments and dependencies"),
    ("C7", "If available, link to developer documentation/manual"),
    ("C8", "Support email for questions"),
]

#: Sections 1-5 must fit six pages and 4000 words, excluding metadata, tables, figures
#: and references.
WORD_LIMIT = 4000

#: The template asks for ca. 100 words. It is a guide, not a hard limit, so an over-long
#: abstract is reported as a warning rather than a failure.
ABSTRACT_TARGET = 100

AFFILIATION = "Institute of One, LISIT Co., Ltd., Tokyo 150-0044, Japan"

#: The release the manuscript describes, and its Zenodo DOIs. v0.1.0 crashed on NumPy 2
#: in the CBF/CBV/MTT/TTP path that produces Fig. 4, which is why v0.1.1 exists; a
#: manuscript that still cites v0.1.0 sends a reviewer to software that does not run.
RELEASE = "v0.1.1"
VERSION_DOI = "10.5281/zenodo.22226447"
CONCEPT_DOI = "10.5281/zenodo.20921268"
SUPERSEDED = ("v0.1.0", "10.5281/zenodo.20921269")


def outline(doc) -> list[tuple[str, str]]:
    return [
        (p.style.name, p.text.strip())
        for p in doc.paragraphs
        if p.text.strip() and p.style.name.startswith("Heading")
    ]


def body_words(doc) -> int:
    """Words in sections 1-5: everything between the first numbered heading and the
    Declarations block, excluding headings and figure captions."""
    count = 0
    inside = False
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        if p.style.name.startswith("Heading"):
            if re.match(r"^1\.\s", text):
                inside = True
            elif text in ("Declarations", "References"):
                inside = False
            continue
        if inside and not re.match(r"^Fig\.\s*\d", text):
            count += len(text.split())
    return count


def abstract_words(doc) -> int:
    count = 0
    inside = False
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        if p.style.name.startswith("Heading"):
            inside = text == "Abstract"
            continue
        if inside and not text.startswith("Keywords"):
            count += len(text.split())
    return count


def check() -> tuple[list[str], list[str]]:
    doc = Document(DOCX)
    heads = outline(doc)
    titles = [t for _, t in heads]
    flat = " ".join(p.text for p in doc.paragraphs)

    fatal: list[str] = []
    warn: list[str] = []

    # -- the returned-submission defects -------------------------------------------
    numbered = [t for t in titles if re.match(r"^\d+\.\s", t)]
    if numbered != NUMBERED_SECTIONS:
        fatal.append(
            "numbered sections do not match the template\n"
            f"      expected: {NUMBERED_SECTIONS}\n"
            f"      found:    {numbered}"
        )
    for stale in ("Limitations and future work", "Reproducibility"):
        if any(stale in t for t in titles):
            fatal.append(f"section the editorial office asked us to remove is still present: {stale}")

    for sub in SUBSECTIONS:
        if sub not in titles:
            fatal.append(f"missing subsection heading: {sub}")

    # -- the code metadata table ----------------------------------------------------
    tables = [t for t in doc.tables if t.rows and t.rows[0].cells[0].text.strip() == "Nr"]
    if not tables:
        fatal.append("the code metadata table is missing, or its first column is not headed 'Nr'")
    else:
        table = tables[0]
        header = [c.text.strip() for c in table.rows[0].cells]
        if header != ["Nr", "Code metadata description", "Metadata"]:
            fatal.append(f"code metadata column headings are {header}")
        got = [(r.cells[0].text.strip(), r.cells[1].text.strip()) for r in table.rows[1:]]
        if got != CODE_METADATA:
            fatal.append(
                "code metadata rows do not match the template\n"
                f"      expected: {[n for n, _ in CODE_METADATA]}\n"
                f"      found:    {[n for n, _ in got]}"
            )
        else:
            values = {r.cells[0].text.strip(): r.cells[2].text.strip() for r in table.rows[1:]}
            if "github.com" not in values["C2"]:
                fatal.append(
                    "C2 is not a GitHub repository. The template states that a GitHub "
                    "repository is mandatory and that the paper will not proceed otherwise.\n"
                    f"      found: {values['C2']}"
                )
            if "@" not in values["C8"]:
                fatal.append(f"C8 is not an email address: {values['C8']}")

    # -- things that must survive the restructuring ---------------------------------
    # Removing section 6 must not take the archive or the run instructions with it.
    for survivor, where in (
        (VERSION_DOI, "the archived version DOI"),
        (CONCEPT_DOI, "the concept DOI"),
        ("run_synthetic_demo.py", "the command that regenerates Fig. 1"),
        ("run_snr_sweep.py", "the command that regenerates Fig. 2"),
        ("make_a_lut_figures.py", "the command that regenerates Fig. 3"),
        ("make_unitobrain_figure.py", "the command that regenerates Fig. 4"),
        ("pytest tests", "the command that runs the test suite"),
        ("must not be used as a medical device", "the medical-device disclaimer"),
    ):
        if survivor not in flat:
            fatal.append(f"lost when sections 5 and 6 were removed: {where} ({survivor})")

    if AFFILIATION not in flat:
        fatal.append(f"affiliation is not the canonical string: {AFFILIATION}")

    if RELEASE not in flat:
        fatal.append(f"the manuscript does not name the release it describes: {RELEASE}")
    for stale in SUPERSEDED:
        if stale in flat:
            fatal.append(
                f"the manuscript still carries the superseded {stale}; that release fails "
                "on NumPy 2 in the path that produces Fig. 4"
            )

    # -- length ----------------------------------------------------------------------
    words = body_words(doc)
    if words > WORD_LIMIT:
        fatal.append(f"sections 1-5 are {words} words, over the {WORD_LIMIT}-word limit")

    abstract = abstract_words(doc)
    if abstract > ABSTRACT_TARGET * 1.5:
        warn.append(
            f"the abstract is {abstract} words; the template asks for ca. {ABSTRACT_TARGET}"
        )

    print(f"sections 1-5: {words} words (limit {WORD_LIMIT})")
    print(f"abstract:     {abstract} words (template asks ca. {ABSTRACT_TARGET})")
    print("outline:")
    for style, title in heads:
        print(f"  {'  ' * (int(style[-1]) - 1)}{title}")
    print()

    return fatal, warn


def main() -> int:
    fatal, warn = check()
    for item in fatal:
        print(f"  FAIL  {item}")
    for item in warn:
        print(f"  WARN  {item}")
    if not fatal:
        print("conforms to the SoftwareX Original Software Publication template")
    return 1 if fatal else 0


if __name__ == "__main__":
    sys.exit(main())
